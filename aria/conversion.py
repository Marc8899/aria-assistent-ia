# Leer y convertir documentos. Para formatos de Office usa el propio Office instalado (vía COM),
# que da resultados perfectos; para lo demás, librerías de Python.
import csv
import os
import re
from contextlib import contextmanager
from pathlib import Path

from aria.archivos import comprobar_escritura, ruta_libre
from aria.texto import recortar

FORMATOS_SOPORTADOS = (
    "Admite: Word (docx/doc/rtf/odt) -> pdf, docx, txt, md, html; PDF -> docx, txt, md; "
    "Excel (xlsx/xls/csv) -> pdf, xlsx, csv; PowerPoint (pptx/ppt) -> pdf; txt/md -> docx, pdf; "
    "imágenes (png, jpg, webp, bmp, gif, tiff, ico) entre sí y a pdf."
)
IMAGENES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff", ".ico"}
WORD = {".docx", ".doc", ".rtf", ".odt", ".dotx"}
EXCEL = {".xlsx", ".xlsm", ".xls", ".ods"}
POWERPOINT = {".pptx", ".ppt", ".odp"}
TEXTO = {".txt", ".md", ".markdown", ".json", ".py", ".js", ".html", ".htm", ".xml", ".log", ".ini", ".yaml",
         ".yml", ".css", ".sql", ".bat", ".ps1", ".toml"}
FORMATO_WORD = {"pdf": 17, "docx": 16, "txt": 7, "html": 10, "rtf": 6, "odt": 23}


@contextmanager
def _office(aplicacion: str):
    """Abre una instancia invisible y separada de Word/Excel/PowerPoint y la cierra al terminar."""
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx(f"{aplicacion}.Application")
    try:
        if aplicacion != "PowerPoint":
            app.Visible = False
        app.DisplayAlerts = 0 if aplicacion != "Excel" else False
        yield app
    finally:
        try:
            app.Quit()
        except Exception:
            pass


def _salida(origen: Path, formato: str) -> Path:
    destino = ruta_libre(origen.with_suffix("." + formato))
    comprobar_escritura(destino)
    return destino


# ---------- Markdown sencillo -> Word ----------

def _runs_con_negrita(parrafo, texto: str):
    for i, trozo in enumerate(re.split(r"\*\*(.+?)\*\*", texto)):
        if trozo:
            parrafo.add_run(trozo).bold = bool(i % 2)


def markdown_a_docx(contenido: str, destino: Path, titulo: str = ""):
    """Admite: # títulos, - viñetas, 1. listas numeradas, **negrita**, tablas con | y párrafos."""
    from docx import Document

    doc = Document()
    if titulo:
        doc.add_heading(titulo, level=0)
    lineas = contenido.splitlines()
    i = 0
    while i < len(lineas):
        linea = lineas[i].rstrip()
        if linea.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                celdas = [c.strip() for c in lineas[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in celdas):  # salta la fila |---|---|
                    filas.append(celdas)
                i += 1
            if filas:
                tabla = doc.add_table(rows=len(filas), cols=max(len(f) for f in filas))
                tabla.style = "Table Grid"
                for r, fila in enumerate(filas):
                    for c, valor in enumerate(fila):
                        celda = tabla.cell(r, c)
                        celda.text = ""
                        _runs_con_negrita(celda.paragraphs[0], valor)
                        if r == 0:
                            for run in celda.paragraphs[0].runs:
                                run.bold = True
            continue
        titulo_md = re.match(r"(#{1,4})\s+(.*)", linea)
        if titulo_md:
            doc.add_heading(titulo_md.group(2), level=len(titulo_md.group(1)))
        elif re.match(r"\s*[-*•]\s+", linea):
            _runs_con_negrita(doc.add_paragraph(style="List Bullet"), re.sub(r"\s*[-*•]\s+", "", linea, count=1))
        elif re.match(r"\s*\d+[.)]\s+", linea):
            _runs_con_negrita(doc.add_paragraph(style="List Number"), re.sub(r"\s*\d+[.)]\s+", "", linea, count=1))
        elif linea.strip():
            _runs_con_negrita(doc.add_paragraph(), linea)
        i += 1
    doc.save(destino)


# ---------- Leer ----------

def _rango_paginas(texto: str, total: int) -> list[int]:
    if not texto.strip():
        return list(range(total))
    paginas = set()
    for parte in texto.split(","):
        if "-" in parte:
            a, b = parte.split("-", 1)
            paginas.update(range(int(a) - 1, min(int(b), total)))
        elif parte.strip().isdigit():
            paginas.add(int(parte) - 1)
    return sorted(p for p in paginas if 0 <= p < total)


def leer_texto(ruta: Path, paginas: str = "", maximo: int = 9000) -> str:
    sufijo = ruta.suffix.lower()
    if sufijo == ".pdf":
        from pypdf import PdfReader
        lector = PdfReader(str(ruta))
        total = len(lector.pages)
        partes = [f"--- página {p + 1} ---\n{lector.pages[p].extract_text() or ''}" for p in _rango_paginas(paginas, total)]
        texto = "\n".join(partes)
        if not re.search(r"\w{3,}", texto):
            return f"PDF de {total} páginas sin texto seleccionable (parece escaneado; no tengo OCR todavía)."
        return f"PDF de {total} páginas.\n" + recortar(texto, maximo)
    if sufijo == ".docx":
        from docx import Document
        doc = Document(str(ruta))
        partes = []
        for p in doc.paragraphs:
            if p.text.strip():
                estilo = p.style.name.lower() if p.style is not None else ""
                prefijo = "# " if "heading" in estilo or "título" in estilo else ""
                partes.append(prefijo + p.text)
        for n, tabla in enumerate(doc.tables, 1):
            partes.append(f"\n[Tabla {n}]")
            partes += [" | ".join(c.text.strip() for c in fila.cells) for fila in tabla.rows]
        return recortar("\n".join(partes), maximo) or "(documento vacío)"
    if sufijo in WORD:  # .doc, .rtf, .odt: Word lo pasa a texto
        temporal = ruta_libre(Path.home() / "AppData/Local/Temp" / (ruta.stem + "_aria.txt"))
        with _office("Word") as word:
            doc = word.Documents.Open(str(ruta), ReadOnly=True, ConfirmConversions=False)
            doc.SaveAs2(str(temporal), FileFormat=7)
            doc.Close(False)
        texto = temporal.read_text(encoding="utf-16", errors="replace")
        temporal.unlink(missing_ok=True)
        return recortar(texto, maximo)
    if sufijo in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        valores = load_workbook(str(ruta), read_only=True, data_only=True)
        formulas = load_workbook(str(ruta), read_only=True)  # para fórmulas aún sin calcular por Excel
        partes = []
        for hoja, hoja_f in zip(valores.worksheets, formulas.worksheets):
            partes.append(f"--- hoja '{hoja.title}' ({hoja.max_row} filas x {hoja.max_column} columnas) ---")
            filas_f = hoja_f.iter_rows(max_row=80, max_col=25, values_only=True)
            for fila, fila_f in zip(hoja.iter_rows(max_row=80, max_col=25, values_only=True), filas_f):
                celdas = [f if v is None and isinstance(f, str) and f.startswith("=") else v
                          for v, f in zip(fila, fila_f)]
                while celdas and celdas[-1] is None:
                    celdas.pop()
                if celdas:
                    partes.append(" | ".join("" if v is None else str(v) for v in celdas))
        valores.close()
        formulas.close()
        return recortar("\n".join(partes), maximo)
    if sufijo in (".xls", ".ods"):
        return leer_texto(convertir(ruta, "xlsx"), paginas, maximo)
    if sufijo == ".csv":
        with open(ruta, encoding="utf-8-sig", errors="replace", newline="") as f:
            muestra = f.read(4096)
            f.seek(0)
            dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t") if muestra else csv.excel
            filas = [" | ".join(fila) for _, fila in zip(range(100), csv.reader(f, dialecto))]
        return recortar("\n".join(filas), maximo)
    if sufijo in (".pptx",):
        from pptx import Presentation
        presentacion = Presentation(str(ruta))
        partes = []
        for n, diapositiva in enumerate(presentacion.slides, 1):
            textos = [s.text_frame.text for s in diapositiva.shapes if s.has_text_frame and s.text_frame.text.strip()]
            partes.append(f"--- diapositiva {n} ---\n" + "\n".join(textos))
        return recortar("\n".join(partes), maximo)
    if sufijo in TEXTO or sufijo == "":
        return recortar(ruta.read_text(encoding="utf-8", errors="replace"), maximo)
    raise ValueError(f"No sé leer archivos {sufijo}.")


# ---------- Convertir ----------

def convertir(origen: Path, formato: str) -> Path:
    sufijo = origen.suffix.lower()
    formato = {"jpeg": "jpg", "tif": "tiff", "markdown": "md", "texto": "txt", "word": "docx",
               "excel": "xlsx"}.get(formato, formato)
    if sufijo == "." + formato:
        raise ValueError("ya está en ese formato")

    if sufijo in IMAGENES:
        from PIL import Image
        destino = _salida(origen, formato)
        with Image.open(origen) as imagen:
            if formato in ("jpg", "pdf", "bmp") and imagen.mode in ("RGBA", "LA", "P"):
                fondo = Image.new("RGB", imagen.size, "white")
                fondo.paste(imagen.convert("RGBA"), mask=imagen.convert("RGBA").split()[-1])
                imagen = fondo
            imagen.save(destino, **({"quality": 92} if formato in ("jpg", "webp") else {}))
        return destino

    if sufijo == ".pdf" and formato == "docx":
        # pdf2docx conserva bien el formato y no depende de Word (Word muestra un aviso que bloquea).
        from pdf2docx import Converter
        destino = _salida(origen, "docx")
        conversor = Converter(str(origen))
        try:
            conversor.convert(str(destino))
        finally:
            conversor.close()
        return destino

    if sufijo in WORD:
        if formato == "md":
            texto = leer_texto(origen, maximo=10_000_000)
            destino = _salida(origen, "md")
            destino.write_text(texto, encoding="utf-8")
            return destino
        if formato not in FORMATO_WORD:
            raise ValueError(f"Word no exporta a {formato}")
        destino = _salida(origen, formato)
        with _office("Word") as word:
            doc = word.Documents.Open(str(origen), ReadOnly=True, ConfirmConversions=False, AddToRecentFiles=False)
            doc.SaveAs2(str(destino), FileFormat=FORMATO_WORD[formato])
            doc.Close(False)
        return destino

    if sufijo == ".pdf" and formato in ("txt", "md"):
        destino = _salida(origen, formato)
        destino.write_text(leer_texto(origen, maximo=10_000_000), encoding="utf-8")
        return destino

    if sufijo in (".txt", ".md", ".markdown") and formato in ("docx", "pdf"):
        contenido = origen.read_text(encoding="utf-8", errors="replace")
        if formato == "docx":
            docx = _salida(origen, "docx")
            markdown_a_docx(contenido, docx)
            return docx
        # A PDF: un .docx temporal que Word exporta y luego se borra.
        temporal = ruta_libre(Path(os.environ.get("TEMP", Path.home())) / f"{origen.stem}.docx")
        markdown_a_docx(contenido, temporal)
        destino = _salida(origen, "pdf")
        try:
            with _office("Word") as word:
                doc = word.Documents.Open(str(temporal), ReadOnly=True, AddToRecentFiles=False)
                doc.SaveAs2(str(destino), FileFormat=17)
                doc.Close(False)
        finally:
            temporal.unlink(missing_ok=True)
        return destino

    if sufijo == ".csv" and formato == "xlsx":
        from openpyxl import Workbook
        destino = _salida(origen, "xlsx")
        libro = Workbook()
        hoja = libro.active
        with open(origen, encoding="utf-8-sig", errors="replace", newline="") as f:
            muestra = f.read(4096)
            f.seek(0)
            dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t") if muestra else csv.excel
            for fila in csv.reader(f, dialecto):
                hoja.append(fila)
        libro.save(destino)
        return destino

    if sufijo in EXCEL or sufijo == ".csv":
        destino = _salida(origen, formato)
        with _office("Excel") as excel:
            libro = excel.Workbooks.Open(str(origen), ReadOnly=True)
            if formato == "pdf":
                libro.ExportAsFixedFormat(0, str(destino))
            elif formato in ("xlsx", "csv"):
                libro.SaveAs(str(destino), FileFormat={"xlsx": 51, "csv": 62}[formato])  # 62 = CSV UTF-8
            else:
                raise ValueError(f"Excel no exporta a {formato}")
            libro.Close(False)
        return destino

    if sufijo in POWERPOINT and formato == "pdf":
        destino = _salida(origen, "pdf")
        with _office("PowerPoint") as ppt:
            presentacion = ppt.Presentations.Open(str(origen), ReadOnly=True, WithWindow=False)
            presentacion.SaveAs(str(destino), 32)
            presentacion.Close()
        return destino

    raise ValueError(f"no sé convertir de {sufijo} a {formato}. {FORMATOS_SOPORTADOS}")
